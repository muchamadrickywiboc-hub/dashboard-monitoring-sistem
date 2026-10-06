# Bangun spreadsheet Summary (format yang dibaca Kode.gs) dari agregat file mentah + Juli dari Summary lama
import json, glob, os, re, sys, collections, datetime, calendar
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
AGG, OLD, OUTX, OUTJ = sys.argv[1:5]
BLN = ['Januari','Februari','Maret','April','Mei','Juni','Juli','Agustus','September','Oktober','November','Desember']
B3 = [b[:3] for b in BLN]
REG = ['SUMBAGUT','SUMBAGSEL','JABALINUS','KALIMANTAN','SULAWESI','MALUPA']
RN = {r: r.title() for r in REG}
KAT = {'DB': 'Driving Behaviour', 'DD': 'Driver Discipline', 'FM': 'Fatigue Management'}
PAR = [  # label, kategori, kode case di data mentah
  ('Over Speed','DB',['OVERSPEED','OVER SPEED']), ('Harsh Turn / Cornering','DB',['HARSH TURN']),
  ('Harsh Braking','DB',['HARSH BREAKING','HARSH BRAKING']), ('Harsh Acceleration','DB',['HARSH ACCELERATION']),
  ('Driving > 4 Hours','DB',['DRIVING > 4 HOURS']),
  ('Black Zone','DD',['BLACKZONE','BLACK ZONE']), ('Idling','DD',['IDLE','IDLING']),
  ('Menggunakan Telepon','DD',['PHONE DETECTION']), ('Merokok / Vape','DD',['SMOKING DETECTION']),
  ('Microsleep','FM',['DRIVER FATIGUE']), ('Menguap','FM',['YAWNING DETECTION']),
  ('Distraction Pengemudi','LAIN',['DRIVER DISTRACTION']), ('Camera Covering','LAIN',['CAMERA COVERING ALARM']),
]
CASE = {c: (l, k) for l, k, cs in PAR for c in cs}
norm_case = lambda c: re.sub(r'\s*>\s*', ' > ', c).strip()
def reg_of(area):
    for r in REG:
        if area.startswith(r): return r
    return None

# ---------- baca agregat ----------
C = collections.Counter()     # (m, R, lok, label)
DAY = collections.Counter()   # (m, d, R, kat)
AMT = collections.Counter()   # (m, R, name, lok, label)
excl = collections.Counter(); drop = collections.Counter(); noreg = 0; amt_blank = 0; amt_total = 0
months_raw = set(); files = []
lok_reg = collections.defaultdict(collections.Counter)
raw = []
for f in sorted(glob.glob(os.path.join(AGG, '*.json'))):
    d = json.load(open(f)); b = os.path.basename(f)[:-5]
    fm = 7 if b.startswith('rawagu_') else BLN.index(b.split('_', 1)[1])
    months_raw.add(fm); files.append((b, fm, d['info']['rows']))
    raw.append((fm, d))
    for m, area, lok, case, n in d['C']:
        r = reg_of(area)
        if r and lok: lok_reg[lok][r] += n
for fm, d in raw:
    def mm(m, case, n):
        if m == -1: return fm
        if m != fm: drop[(fm, m)] += n; return None
        return m
    for m, area, lok, case, n in d['C']:
        m2 = mm(m, case, n)
        if m2 is None: continue
        case = norm_case(case)
        if case not in CASE: excl[case] += n; continue
        r = reg_of(area) or (lok_reg[lok].most_common(1)[0][0] if lok_reg.get(lok) else None)
        if not r: noreg += n; r = '?'
        C[(m2, r, lok or '(TANPA LOKASI)', CASE[case][0])] += n
    for m, dd, area, case, n in d['DAY']:
        if m == -1 or m != fm: continue
        case = norm_case(case)
        if case not in CASE or CASE[case][1] == 'LAIN': continue
        r = reg_of(area)
        if r: DAY[(m, dd, r, CASE[case][1])] += n
    for m, area, lok, name, case, n in d['AMT']:
        if m == -1: m = fm
        if m != fm: continue
        case = norm_case(case)
        if case not in CASE or CASE[case][1] == 'LAIN': continue
        amt_total += n
        nm = re.sub(r'\s+', ' ', re.sub(r'\[.*?\]|\(.*?\)', ' ', name)).strip(' .-')
        if not nm or nm in ('0', 'NONE', '(BLANK)', 'BLANK', 'N/A', '#N/A', 'TIDAK ADA', 'KOSONG'): amt_blank += n; continue
        r = reg_of(area) or (lok_reg[lok].most_common(1)[0][0] if lok_reg.get(lok) else None)
        if not r: amt_blank += n; continue
        AMT[(m, r, nm, lok, CASE[case][0])] += n

# ---------- Juli dari Summary lama ----------
old = json.load(open(OLD))
JUL = 6
jul_par, jul_reg, jul_lok = {}, {}, {}
hdr = None
for row in old['Parameter_Bulanan']:
    a = str(row[0]).strip()
    if a == 'Parameter': hdr = [str(x).strip() for x in row]; continue
    if hdr and 'Juli' in hdr and a and not a.upper().startswith('TOTAL') and row[hdr.index('Juli')] not in ('', None):
        jul_par[a] = row[hdr.index('Juli')]
hdr = None
for row in old['Rekap_Regional']:
    a = str(row[0]).strip()
    if a == 'Regional': hdr = [str(x).strip() for x in row]; continue
    if hdr and 'Juli' in hdr and a.upper() in REG: jul_reg[a.upper()] = row[hdr.index('Juli')]
hdr = None
for row in old['Rekap_Lokasi_Agu']:
    a = str(row[0]).strip()
    if a == 'Lokasi': hdr = [str(x).strip() for x in row]; continue
    if hdr and a: jul_lok[a.upper()] = (row[1].upper(), row[hdr.index('Total Jul')] or 0)
has_jul = JUL not in months_raw and bool(jul_par)
MON = sorted(months_raw | ({JUL} if has_jul else set()))
LASTM = MON[-1]
lbl = lambda m: BLN[m] + ('*' if m == 7 else '')

# ---------- turunan ----------
par = collections.defaultdict(lambda: [None] * 12)
for (m, r, lok, l), n in C.items():
    par[l][m] = (par[l][m] or 0) + n
for l, k, _ in PAR:
    for m in months_raw: par[l][m] = par[l][m] or 0
    if has_jul: par[l][JUL] = jul_par.get(l, 0) or 0
katOf = {l: k for l, k, _ in PAR}
regkat = collections.defaultdict(lambda: {'DB': 0, 'DD': 0, 'FM': 0})   # (m,R)
lokkat = collections.defaultdict(lambda: {'DB': 0, 'DD': 0, 'FM': 0})   # (m,R,lok)
for (m, r, lok, l), n in C.items():
    k = katOf[l]
    if k == 'LAIN' or r == '?': continue
    regkat[(m, r)][k] += n; lokkat[(m, r, lok)][k] += n
regtot = {(m, r): sum(regkat[(m, r)].values()) for m in months_raw for r in REG}
if has_jul:
    for r in REG: regtot[(JUL, r)] = jul_reg.get(r)
lokset = collections.defaultdict(set)
for (m, r, lok) in lokkat: lokset[r].add(lok)
if has_jul:
    for lok, (rn, t) in jul_lok.items():
        if t: lokset[rn.upper()].add(lok)
prevm = lambda m: MON[MON.index(m) - 1] if MON.index(m) > 0 else None
def lok_total(m, r, lok):
    if m == JUL and has_jul:
        x = jul_lok.get(lok); return x[1] if x and x[0].upper() == r else None
    k = lokkat.get((m, r, lok)); return sum(k.values()) if k else 0

wb = openpyxl.Workbook(); wb.remove(wb.active)
J = {}
BOLD = Font(bold=True); HFILL = PatternFill('solid', fgColor='4A7A53'); HFONT = Font(bold=True, color='FFFFFF')
def sheet(name, rows, widths=None, heads=()):
    ws = wb.create_sheet(name)
    for row in rows: ws.append(row)
    for i, row in enumerate(rows, 1):
        if row and isinstance(row[0], str) and (i == 1 or row[0].upper() in heads or row[0].upper() in ('PARAMETER', 'REGIONAL', 'LOKASI', 'BULAN', 'TANGGAL')):
            for c in ws[i]:
                if i == 1: c.font = Font(bold=True, size=13)
                elif c.value is not None: c.font = HFONT; c.fill = HFILL
        elif row and isinstance(row[0], str) and (row[0].upper().startswith('TOTAL') or row[0].upper() in KAT.values() or row[0].upper() in [v.upper() for v in KAT.values()]):
            for c in ws[i]: c.font = BOLD
    for j, w in enumerate(widths or [], 1): ws.column_dimensions[openpyxl.utils.get_column_letter(j)].width = w
    for r in ws.iter_rows():
        for c in r:
            if isinstance(c.value, float) and abs(c.value) < 1000 and c.column_letter != 'A' and 'Kenaikan' in str(ws.cell(row=1, column=1).value or '') : pass
    J[name] = [[('' if v is None else v) for v in row] for row in rows]
    return ws
def growth(a, b):
    if a is None or b is None: return ''
    if b == 0: return 'baru' if a else ''
    return round(a / b - 1, 6)

# Parameter_Bulanan
rng = f'{B3[MON[0]]}–{B3[LASTM]}'
pm_last = prevm(LASTM)
H = ['Parameter'] + [lbl(m) for m in MON] + [f'Total {rng}', f'Kenaikan {B3[pm_last]}→{B3[LASTM]}']
rows = [['Parameter per Kategori per Bulan'], ['Dibangun dari file laporan mentah per bulan. * Agustus belum lengkap, lihat sheet Catatan.'], [], H]
tot3 = [0] * 12
for k in ('DB', 'DD', 'FM'):
    rows.append([KAT[k].upper()])
    kt = [0] * 12
    for l, kk, _ in PAR:
        if kk != k: continue
        v = [par[l][m] for m in MON]
        for m in MON: kt[m] += par[l][m] or 0
        rows.append([l] + v + [sum(x or 0 for x in v), growth(par[l][LASTM], par[l][pm_last])])
    for m in MON: tot3[m] += kt[m]
    rows.append([f'TOTAL {KAT[k]}'] + [kt[m] for m in MON] + [sum(kt[m] for m in MON), growth(kt[LASTM], kt[pm_last])])
    rows.append([])
rows.append(['TOTAL 3 KATEGORI'] + [tot3[m] for m in MON] + [sum(tot3[m] for m in MON), growth(tot3[LASTM], tot3[pm_last])])
rows += [[], [], ['PARAMETER LAIN (di luar 3 kategori, tidak masuk total)'], H]
for l, kk, _ in PAR:
    if kk == 'LAIN':
        v = [par[l][m] for m in MON]
        rows.append([l] + v + [sum(x or 0 for x in v), growth(par[l][LASTM], par[l][pm_last])])
sheet('Parameter_Bulanan', rows, [30] + [12] * (len(MON) + 2))

# Rekap_Regional
rows = [['Rekap per Regional'], ['Satu tabel per bulan (kategori per regional), lalu total 3 kategori per regional per bulan.'], []]
nlok = {r: len(lokset[r]) for r in REG}
for m in [x for x in MON if x in months_raw]:
    p = prevm(m)
    rows.append([f'Rekap per Regional – {BLN[m]} 2026' + (f' vs {BLN[p]} 2026' if p is not None else '')])
    rows.append(['Regional', 'Driving Behaviour', 'Driver Discipline', 'Fatigue Management', f'Total {BLN[m]}'] + ([f'Total {BLN[p]}', 'Kenaikan'] if p is not None else []) + ['Jumlah Lokasi'])
    T = [0, 0, 0, 0, 0]
    for r in REG:
        k = regkat[(m, r)]; t = regtot[(m, r)]; pt = regtot.get((p, r)) if p is not None else None
        rows.append([RN[r], k['DB'], k['DD'], k['FM'], t] + ([pt, growth(t, pt)] if p is not None else []) + [nlok[r]])
        T[0] += k['DB']; T[1] += k['DD']; T[2] += k['FM']; T[3] += t; T[4] += pt or 0
    rows.append(['TOTAL', T[0], T[1], T[2], T[3]] + ([T[4], growth(T[3], T[4])] if p is not None else []) + [sum(nlok.values())])
    rows.append([])
rows.append(['Total event 3 kategori per regional per bulan'])
rows.append(['Regional'] + [lbl(m) for m in MON])
for r in REG: rows.append([RN[r]] + [regtot.get((m, r)) for m in MON])
sheet('Rekap_Regional', rows, [22, 16, 16, 18, 14, 14, 11, 13])

# Rekap_Lokasi_<Bln>
for m in [x for x in MON if x in months_raw]:
    p = prevm(m)
    rows = [[f'Rekap per Lokasi – {BLN[m]} 2026 (urut event terbanyak)'], ['* Agustus belum lengkap, lihat Catatan' if m == 7 else ''], []]
    rows.append(['Lokasi', 'Regional', 'Driving Behaviour', 'Driver Discipline', 'Fatigue Management', 'Total'] + ([f'Total {B3[p]}', 'Kenaikan'] if p is not None else []))
    items = []
    keys = {(r, lok) for (mm_, r, lok) in lokkat if mm_ == m}
    if p is not None:
        keys |= {(r, lok) for (mm_, r, lok) in lokkat if mm_ == p}
        if p == JUL and has_jul: keys |= {(rn.upper(), lok) for lok, (rn, t) in jul_lok.items() if t}
    for r, lok in keys:
        k = lokkat.get((m, r, lok), {'DB': 0, 'DD': 0, 'FM': 0}); t = sum(k.values())
        pt = lok_total(p, r, lok) if p is not None else None
        if not t and not pt: continue
        items.append([lok, RN[r], k['DB'], k['DD'], k['FM'], t] + ([pt, growth(t, pt)] if p is not None else []))
    items.sort(key=lambda x: (-x[5], x[0]))
    sheet(f'Rekap_Lokasi_{B3[m]}', rows + items, [24, 13, 16, 16, 18, 11, 11, 11])

# Rekap_Lokasi_YTD (bulan dengan data lokasi per kategori)
rawM = sorted(months_raw)
rows = [[f'Rekap per Lokasi – {", ".join(B3[m] for m in rawM)} 2026'], ['Juli tidak ikut karena file mentah Juli belum ada (rincian lokasi Juli tidak tersedia). Rata-rata/bulan = total ÷ bulan terdata.' if has_jul else 'Rata-rata/bulan = total ÷ bulan terdata.'], []]
rows.append(['Lokasi', 'Regional', 'Driving Behaviour', 'Driver Discipline', 'Fatigue Management', 'Total', 'Rata-rata/bulan', 'Bulan Terdata'])
agg = collections.defaultdict(lambda: [0, 0, 0, set()])
for (m, r, lok), k in lokkat.items():
    a = agg[(r, lok)]; a[0] += k['DB']; a[1] += k['DD']; a[2] += k['FM']
    if sum(k.values()): a[3].add(m)
items = []
for (r, lok), a in agg.items():
    t = a[0] + a[1] + a[2]
    if not t: continue
    items.append([lok, RN[r], a[0], a[1], a[2], t, round(t / max(len(a[3]), 1), 2), len(a[3])])
items.sort(key=lambda x: (-x[5], x[0]))
sheet('Rekap_Lokasi_YTD', rows + items, [24, 13, 16, 16, 18, 11, 14, 13])

# Top_AMT: 10 teratas per regional per bulan + 10 teratas per regional seluruh periode (Bulan = YTD)
am = collections.defaultdict(lambda: {'DB': 0, 'DD': 0, 'FM': 0, 'par': collections.Counter(), 'lok': collections.Counter()})
amy = collections.defaultdict(lambda: {'DB': 0, 'DD': 0, 'FM': 0, 'par': collections.Counter(), 'lok': collections.Counter()})
for (m, r, nm, lok, l), n in AMT.items():
    for o in (am[(m, r, nm)], amy[(r, nm)]):
        o[katOf[l]] += n; o['par'][l] += n; o['lok'][lok] += n
def toprows(src, bulan_of):
    out = []
    by = collections.defaultdict(list)
    for key, o in src.items():
        g = key[:-1]; by[g].append((key, o))
    for g in sorted(by, key=lambda g: tuple(str(x) for x in g)):
        lst = sorted(by[g], key=lambda x: -(x[1]['DB'] + x[1]['DD'] + x[1]['FM']))[:10]
        for key, o in lst:
            r = key[-2]
            out.append([bulan_of(key), key[-1], RN.get(r, ''), o['lok'].most_common(1)[0][0], o['DB'], o['DD'], o['FM'], o['DB'] + o['DD'] + o['FM'], o['par'].most_common(1)[0][0]])
    return out
rows = [['Top AMT per Regional (10 teratas per regional per bulan, dan per regional seluruh periode = YTD)'],
        ['Nama AMT dari kolom AMT 1. Event tanpa nama AMT tidak ikut dihitung. Juli belum ada (file mentah Juli belum ada).' if has_jul else 'Nama AMT dari kolom AMT 1. Event tanpa nama AMT tidak ikut dihitung.'], []]
rows.append(['Bulan', 'Nama AMT', 'Regional', 'Lokasi', 'Driving Behaviour', 'Driver Discipline', 'Fatigue Management', 'Total', 'Pelanggaran Terbanyak'])
tm = toprows(am, lambda k: BLN[k[0]]); tm.sort(key=lambda x: (BLN.index(x[0]), REG.index(x[2].upper()) if x[2] else 9, -x[7]))
ty = toprows(amy, lambda k: 'YTD'); ty.sort(key=lambda x: (REG.index(x[2].upper()) if x[2] else 9, -x[7]))
sheet('Top_AMT', rows + tm + ty, [11, 30, 13, 22, 16, 16, 18, 10, 24])

# Harian
rows = [['Event per Hari per Regional (3 kategori)'], ['Dari kolom Tanggal file mentah. Juli tidak ada.' if has_jul else 'Dari kolom Tanggal file mentah.'], []]
rows.append(['Tanggal', 'Regional', 'Driving Behaviour', 'Driver Discipline', 'Fatigue Management', 'Total'])
dk = collections.defaultdict(lambda: {'DB': 0, 'DD': 0, 'FM': 0})
for (m, dd, r, k), n in DAY.items(): dk[(m, dd, r)][k] += n
for (m, dd, r) in sorted(dk, key=lambda x: (x[0], x[1], REG.index(x[2]))):
    k = dk[(m, dd, r)]
    rows.append([datetime.date(2026, m + 1, dd), RN[r], k['DB'], k['DD'], k['FM'], sum(k.values())])
ws = sheet('Harian', rows, [12, 13, 16, 16, 18, 10])
for c in ws['A'][4:]: c.number_format = 'yyyy-mm-dd'
J['Harian'] = [[(v.isoformat() if isinstance(v, datetime.date) else v) for v in row] for row in J['Harian']]

# Catatan
lastday = collections.defaultdict(int)
for (m, dd, r, k), n in DAY.items(): lastday[(m, r)] = max(lastday[(m, r)], dd)
cat = ['PEMETAAN', 
  'Driving Behaviour: Over Speed=OVERSPEED, Harsh Turn/Cornering=HARSH TURN, Harsh Braking=HARSH BREAKING, Harsh Acceleration=HARSH ACCELERATION, Driving > 4 Hours=DRIVING > 4 HOURS (dua penulisan digabung, dipindah ke Driving Behaviour).',
  'Driver Discipline: Black Zone=BLACKZONE, Idling=IDLE, Menggunakan Telepon=PHONE DETECTION, Merokok/Vape=SMOKING DETECTION.',
  'Fatigue Management: Microsleep=DRIVER FATIGUE (asumsi, mohon konfirmasi), Menguap=YAWNING DETECTION.',
  'Parameter lain: Distraction Pengemudi=DRIVER DISTRACTION, Camera Covering=CAMERA COVERING ALARM.',
  'Tidak dimasukkan: ' + ', '.join(f'{k} ({v:,})'.replace(',', '.') for k, v in excl.most_common()) + '.',
  'Regional: SUMBAGUT I + II digabung jadi Sumbagut. Bulan dihitung dari kolom Tanggal.',
  'TOP AMT: nama dari kolom AMT 1 (pengemudi). ' + f'{amt_blank:,}'.replace(',', '.') + ' dari ' + f'{amt_total:,}'.replace(',', '.') + ' event 3 kategori tidak punya nama AMT dan tidak ikut peringkat.',
  '', 'KETERBATASAN DATA']
lim = []
if has_jul: lim.append('File mentah Juli belum ada. Angka Juli (parameter dan total regional) diambil dari Summary sebelumnya; rincian kategori per regional/lokasi, harian, dan AMT untuk Juli belum tersedia.')
miss = [RN[r] for r in REG if not regtot.get((7, r))]
if 7 in months_raw:
    s = 'Agustus belum lengkap: ' + (', '.join(miss) + ' tidak ada file Agustus' if miss else '')
    cut = [f'{RN[r]} s/d {lastday[(7, r)]} Agu' for r in REG if lastday.get((7, r)) and lastday[(7, r)] < 31]
    if cut: s += ('; ' if miss else '') + ', '.join(cut)
    lim.append(s + '.')
first = {l: min([m for m in months_raw if par[l][m]] or [None], key=lambda x: 99 if x is None else x) for l, k, _ in PAR}
newp = [l for l, k, _ in PAR if first[l] is not None and first[l] > MON[0]]
if newp: lim.append('Jenis event yang baru muncul belakangan (kenaikannya terlihat ekstrem): ' + ', '.join(f'{l} (mulai {BLN[first[l]]})' for l in newp) + '.')
for (fm, m), n in sorted(drop.items()):
    lim.append(f'File {BLN[fm]} memuat {n:,} baris bertanggal {BLN[m]}; baris ini tidak dihitung agar tidak dobel.'.replace(',', '.'))
if noreg: lim.append(f'{noreg:,}'.replace(',', '.') + ' event tanpa Area/regional yang bisa dikenali hanya masuk angka nasional (parameter), tidak masuk regional/lokasi.')
cat += [f'{i}. {s}' for i, s in enumerate(lim, 1)]
sheet('Catatan', [[c] for c in cat], [140])

wb.save(OUTX)
json.dump(J, open(OUTJ, 'w'), default=str)
print('bulan', [BLN[m] for m in MON], 'raw', sorted(months_raw))
print('drop', dict(drop), 'noreg', noreg, 'excl', dict(excl))
print('amt rows', len(tm), len(ty), 'blank', amt_blank, '/', amt_total)
for l, k, _ in PAR: print(f'  {l:24s}', [par[l][m] for m in MON])
print('tot3', [tot3[m] for m in MON])
for r in REG: print(f'  {r:10s}', [regtot.get((m, r)) for m in MON])
