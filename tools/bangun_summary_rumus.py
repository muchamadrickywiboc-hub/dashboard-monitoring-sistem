# Spreadsheet Summary berbasis RUMUS: sheet Data_* (hasil pivot file mentah) + sheet ringkasan dengan SUMIFS/QUERY
import json, glob, os, re, sys, collections, datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter as CL
from openpyxl.comments import Comment
AGG, OUTX = sys.argv[1:3]
BLN = ['Januari','Februari','Maret','April','Mei','Juni','Juli','Agustus','September','Oktober','November','Desember']
B3 = ['Jan','Feb','Mar','Apr','Mei','Jun','Jul','Agu','Sep','Okt','Nov','Des']
REG = ['SUMBAGUT','SUMBAGSEL','JABALINUS','KALIMANTAN','SULAWESI','MALUPA']
RN = {r: r.title() for r in REG}
KATS = ['Driving Behaviour', 'Driver Discipline', 'Fatigue Management']
KAT3 = 'Driving Behaviour|Driver Discipline|Fatigue Management'
PAR = [
  ('Over Speed','Driving Behaviour',['OVERSPEED','OVER SPEED','OVEESPEED']), ('Harsh Turn / Cornering','Driving Behaviour',['HARSH TURN','HHARSH TURN','HASRH TURN']),
  ('Harsh Braking','Driving Behaviour',['HARSH BREAKING','HARSH BRAKING']),
  ('Harsh Acceleration','Driving Behaviour',['HARSH ACCELERATION','HARSH ACCELRATION','HARSH ACCELARATION','HARSH ACCELERATIO','HARSHACCELERATION','HARSH ACCELERATON']),
  ('Driving > 4 Hours','Driving Behaviour',['DRIVING > 4 HOURS','DRIVING >4 HOURS','DRIVING > HOURS','DRIVING 4 HOURS']),
  ('Black Zone','Driver Discipline',['BLACKZONE','BLACK ZONE']), ('Idling','Driver Discipline',['IDLE','IDLING','IDE']),
  ('Menggunakan Telepon','Driver Discipline',['PHONE DETECTION']), ('Merokok / Vape','Driver Discipline',['SMOKING DETECTION']),
  ('Microsleep','Fatigue Management',['DRIVER FATIGUE']), ('Menguap','Fatigue Management',['YAWNING DETECTION']),
  ('Distraction Pengemudi','Parameter lain',['DRIVER DISTRACTION']), ('Camera Covering','Parameter lain',['CAMERA COVERING ALARM','CAMERA-COVERING ALARM','COVERING ALARM','CAMERA COVERING']),
]
EXCL = ['SEAT BELT DETECTION', 'SEATBELT DETECTION', 'ILLEGAL SHUTDOWN', 'REST AREA', 'REST ARE', 'VIDEO LOSS ALARM', 'HIGH SPEED ALARM', 'ABNORMAL STORAGE ALARM', '0.0']
CASE = {c: (l, k) for l, k, cs in PAR for c in cs}
for c in EXCL: CASE[c] = ('', 'Tidak dihitung')
AREA = {'SUMBAGUT I': 'Sumbagut', 'SUMBAGUT II': 'Sumbagut', 'SUMBAGSEL': 'Sumbagsel', 'JABALINUS': 'Jabalinus',
        'KALIMANTAN': 'Kalimantan', 'SULAWESI': 'Sulawesi', 'MALUPA': 'Malupa'}
def reg_of(area):
    a = area.upper()
    if a in AREA: return AREA[a]
    for r in REG:
        if a.startswith(r): return RN[r]
    return None

# ---------- baca agregat ----------
EV = collections.Counter()    # (bulan, tanggal|None, area, lokasi, case)
AM = collections.Counter()    # (bulan, area, lokasi, nama, case)
lokreg = collections.defaultdict(collections.Counter)
areas, cases = set(), collections.Counter()
drop = collections.Counter(); INFO = []
for f in sorted(glob.glob(os.path.join(AGG, '*.json'))):
    d = json.load(open(f)); fm = d['mon'] - 1; INFO.append(d)
    for ds, area, lok, case, n in d['DL']:
        if ds:
            dt = datetime.date.fromisoformat(ds)
            if dt.month - 1 != fm: drop[(fm, dt.month - 1)] += n; continue
        else: dt = None
        EV[(fm + 1, dt, area, lok, case)] += n
        areas.add(area); cases[case] += n
        r = reg_of(area)
        if r and lok: lokreg[lok][r] += n
    for m, area, lok, name, case, n in d['AMT']:
        if m == -1: m = fm
        if m != fm or not name: continue
        if CASE.get(case, ('', ''))[1] not in KATS: continue
        AM[(fm + 1, area, lok, name, case)] += n
unk = [c for c in cases if c not in CASE]
assert not unk, unk
LOKREG = {l: c.most_common(1)[0][0] for l, c in lokreg.items()}
def regrow(area, lok): return reg_of(area) or LOKREG.get(lok) or '?'
MONTHS = sorted({k[0] for k in EV})          # 1-based bulan dengan data mentah
AMONTHS = sorted({k[0] for k in AM})

# ---------- workbook ----------
wb = openpyxl.Workbook(); wb.remove(wb.active)
H1 = Font(bold=True, size=14); HF = Font(bold=True, color='FFFFFF'); HFILL = PatternFill('solid', fgColor='4A7A53')
BOLD = Font(bold=True); NOTE = Font(italic=True, color='666666'); HELP = Font(color='999999', size=9)
TOT = PatternFill('solid', fgColor='EEF4EF'); INPUT = PatternFill('solid', fgColor='FFF7D6')
def head(ws, r, n=None):
    for c in ws[r][: n or ws.max_column]:
        if c.value is not None: c.font = HF; c.fill = HFILL; c.alignment = Alignment(wrap_text=True, vertical='center')
def widths(ws, ws_w):
    for i, w in enumerate(ws_w, 1): ws.column_dimensions[CL(i)].width = w
DE, DA = 'Data_Event', 'Data_AMT'
def cnt(m): return f'COUNTIF({DE}!$A:$A,{m})'
def sumifs(*crit, col='F', sh=DE):
    s = f'SUMIFS({sh}!${col}:${col}'
    for c, v in zip(crit[::2], crit[1::2]): s += f',{sh}!${c}:${c},{v}'
    return s + ')'
def kat3(*crit, sh=DE): return '+'.join(sumifs(*crit, 'I', f'"{k}"', sh=sh) for k in KATS)
# Sheet agregat (satu QUERY besar) supaya SUMIFS di sheet ringkasan hanya membaca tabel kecil
AL, AP, AA, AH = 'Agg_Lokasi', 'Agg_Param', 'Agg_AMT', 'Agg_Harian'
_ML = {'A': 'A', 'D': 'B', 'G': 'C', 'I': 'D'}; _MA = {'A': 'A', 'D': 'B', 'C': 'C', 'G': 'D', 'I': 'E'}
def sl(*crit):   # SUMIFS di Agg_Lokasi dengan huruf kolom Data_Event
    return sumifs(*[(_ML[c] if i % 2 == 0 else c) for i, c in enumerate(crit)], col='E', sh=AL)
def sa(*crit):   # SUMIFS di Agg_AMT dengan huruf kolom Data_AMT
    return sumifs(*[(_MA[c] if i % 2 == 0 else c) for i, c in enumerate(crit)], col='F', sh=AA)
def cntl(m): return f'COUNTIF({AL}!$A:$A,{m})'
def growth(cur, prev): return f'=IF(OR({prev}="",{prev}=0),IF(N({cur})>0,"baru",""),{cur}/{prev}-1)'

# ===== Panduan =====
ws = wb.create_sheet('Panduan')
pand = [
 ('PANDUAN SPREADSHEET SUMMARY DASHBOARD', H1),
 ('Semua angka di sheet ringkasan dihitung dengan RUMUS dari sheet Data_Event dan Data_AMT. Klik sel mana pun untuk melihat rumusnya.', None),
 ('', None),
 ('ALUR DATA', BOLD),
 ('File mentah per regional per bulan (1 baris = 1 event)  →  pivot  →  Data_Event & Data_AMT  →  rumus  →  sheet ringkasan  →  dashboard', None),
 ('Sumber: folder "2026" (subfolder per bulan, isinya file per regional). Daftar file yang dipakai dan kelengkapan tanggalnya ada di sheet Kelengkapan.', None),
 ('', None),
 ('SHEET', BOLD),
 ('Data_Event : jumlah event per Tanggal × Area × Lokasi × Case. Kolom A–F diisi (tempel hasil pivot), kolom G–I (Regional, Parameter, Kategori) terisi otomatis.', None),
 ('Data_AMT   : jumlah event per Bulan × Area × Lokasi × Nama AMT (kolom AMT 1, NIP dalam kurung dibuang) × Case. Kolom A–F diisi, kolom G–I otomatis.', None),
 ('Pemetaan   : kode Case → Parameter & Kategori (termasuk salah ketik di file mentah), Area → Regional, Lokasi → Regional. Tambahkan baris di sini kalau ada Case/Lokasi baru.', None),
 ('Kelengkapan: per bulan & regional: file sumber, jumlah baris, tanggal yang kosong / sangat rendah. Dipakai untuk mengecek data yang kurang.', None),
 ('Agg_*      : tabel ringkas hasil satu QUERY dari Data_Event/Data_AMT (otomatis). Sheet ringkasan membaca tabel ini agar ringan. Jangan diedit.', None),
 ('Cek        : pemeriksaan otomatis (Case/Regional yang belum terpetakan, lokasi yang belum ada di tabel). Semua harus 0 / OK.', None),
 ('Parameter_Bulanan, Rekap_Regional, Rekap_Lokasi_*, Top_AMT, Harian : sheet yang dibaca dashboard. Jangan ganti nama sheet & judul kolomnya.', None),
 ('', None),
 ('CARA UPDATE DATA BULAN BARU (contoh: November)', BOLD),
 ('1. Untuk tiap file regional November (kolom Case, Area, Tanggal, Lokasi, AMT 1): buat Pivot Table, Baris = Tanggal, Area, Lokasi, Case; Nilai = COUNTA dari Case.', None),
 ('   Tampilkan dalam bentuk tabel (tanpa subtotal / total). Salin hasilnya. File yang dipecah (mis. 1–22 dan 23–30) dipivot masing-masing; pastikan tanggalnya tidak tumpang tindih.', None),
 ('2. Di Data_Event, tempel di baris kosong pertama mulai kolom B (Tanggal, Area, Lokasi, Case, Jumlah). Isi kolom A (Bulan) dengan angka 11 untuk semua baris itu.', None),
 ('3. Pivot kedua: Baris = Area, Lokasi, AMT 1, Case; Nilai = COUNTA. Tempel di Data_AMT mulai kolom B, isi kolom A (Bulan) = 11. Baris tanpa nama AMT boleh dibuang.', None),
 ('4. Parameter_Bulanan & total di Rekap_Regional langsung terisi (kolom bulan sudah ada sampai Desember). Sheet Harian juga otomatis.', None),
 ('5. Rekap_Regional: salin satu blok bulan (judul s/d baris TOTAL) ke bawah blok terakhir, ganti judulnya, lalu ganti angka bulan di sel kuning kolom J (11) dan K (10 = bulan pembanding).', None),
 ('   Judul kolom "Total ..." ikut berubah otomatis.', None),
 ('6. Rekap_Lokasi: klik kanan sheet Rekap_Lokasi_Okt → Duplikat → ganti nama jadi Rekap_Lokasi_Nov, lalu ganti sel kuning J1 = 11 dan K1 = 10.', None),
 ('7. Top_AMT: salin 60 baris blok Oktober (6 regional × 10) ke atas blok YTD, lalu ganti angka bulan di kolom kuning K menjadi 11 dan teks bulan di kolom L menjadi November.', None),
 ('8. Buka sheet Cek. Kalau ada Case/Area/Lokasi baru (tanda "?"), tambahkan ke Pemetaan; kalau ada lokasi baru, tambahkan barisnya di Rekap_Lokasi_* (salin baris di atasnya).', None),
 ('9. Dashboard membaca data baru dalam 10 menit (atau menu Dashboard → Hapus cache).', None),
 ('   Alternatif: kirim folder bulan baru ke Claude; Data_Event/Data_AMT dan sheet Kelengkapan dibuat ulang otomatis.', None),
 ('', None),
 ('MENELUSURI ANGKA', BOLD),
 ('Contoh: Idling Agustus di Parameter_Bulanan = SUMIFS(Data_Event Jumlah; Bulan = 8; Parameter = "Idling").', None),
 ('Di Data_Event, filter kolom A = 8 dan H = Idling untuk melihat baris penyusunnya. Tiap baris = jumlah baris di file mentah dengan Tanggal, Area, Lokasi dan Case yang sama,', None),
 ('jadi bisa dicek ulang di file mentah dengan filter yang sama (atau COUNTIFS).', None),
 ('', None),
 ('ATURAN YANG DIPAKAI SAAT MENYUSUN DATA (Jan–Okt)', BOLD),
 ('• Semua sheet yang punya kolom Case, Area dan Tanggal ikut dihitung (mis. "Master Table" dan "Master Table II"); sheet pivot/master data diabaikan.', None),
 ('• Baris yang tanggalnya di luar bulan folder tidak dimasukkan agar tidak dobel. Baris yang persis sama di dua file berbeda (file tumpang tindih) dihitung sekali.', None),
 ('• Data_AMT hanya memuat Case yang masuk 3 kategori dan baris yang punya nama AMT 1.', None),
]
for i, (t, f) in enumerate(pand, 1):
    ws.cell(i, 1, t).font = f or Font()
widths(ws, [160])

# ===== Pemetaan =====
ws = wb.create_sheet('Pemetaan')
ws.append(['Case (kode di file mentah)', 'Parameter', 'Kategori', None, 'Area', 'Regional', None, 'Lokasi', 'Regional'])
mp = [(c, l, k) for l, k, cs in PAR for c in cs] + [(c, '', 'Tidak dihitung') for c in EXCL]
ar = sorted(AREA.items())
lk = sorted(LOKREG.items(), key=lambda x: (REG.index(x[1].upper()), x[0]))
for i in range(max(len(mp), len(ar), len(lk))):
    row = list(mp[i]) if i < len(mp) else [None] * 3
    row += [None] + (list(ar[i]) if i < len(ar) else [None, None])
    row += [None] + (list(lk[i]) if i < len(lk) else [None, None])
    ws.append(row)
head(ws, 1); widths(ws, [26, 24, 20, 3, 14, 13, 3, 24, 13])

# ===== Data_Event =====
ws = wb.create_sheet(DE)
ws.append(['Bulan', 'Tanggal', 'Area', 'Lokasi', 'Case', 'Jumlah',
  '=ARRAYFORMULA(IF(ROW(E:E)=1,"Regional",IF(E:E="","",IFERROR(VLOOKUP(C:C,Pemetaan!E:F,2,FALSE),IFERROR(VLOOKUP(D:D,Pemetaan!H:I,2,FALSE),"?")))))',
  '=ARRAYFORMULA(IF(ROW(E:E)=1,"Parameter",IF(E:E="","",IFERROR(VLOOKUP(E:E,Pemetaan!A:C,2,FALSE),"?"))))',
  '=ARRAYFORMULA(IF(ROW(E:E)=1,"Kategori",IF(E:E="","",IFERROR(VLOOKUP(E:E,Pemetaan!A:C,3,FALSE),"?"))))'])
rows = sorted(EV.items(), key=lambda kv: (kv[0][0], kv[0][1] or datetime.date(1900, 1, 1), kv[0][2], kv[0][3], kv[0][4]))
for (m, dt, area, lok, case), n in rows:
    ws.append([m, dt, area or None, lok or None, case, n])
for c in ws['B'][1:]: c.number_format = 'yyyy-mm-dd'
for c in ws[1][:6]: c.font = HF; c.fill = HFILL
ws.freeze_panes = 'A2'; widths(ws, [7, 12, 13, 24, 24, 9, 12, 22, 19])
n_ev = len(rows)

# ===== Data_AMT =====
ws = wb.create_sheet(DA)
ws.append(['Bulan', 'Area', 'Lokasi', 'Nama AMT', 'Case', 'Jumlah',
  '=ARRAYFORMULA(IF(ROW(E:E)=1,"Regional",IF(E:E="","",IFERROR(VLOOKUP(B:B,Pemetaan!E:F,2,FALSE),IFERROR(VLOOKUP(C:C,Pemetaan!H:I,2,FALSE),"?")))))',
  '=ARRAYFORMULA(IF(ROW(E:E)=1,"Parameter",IF(E:E="","",IFERROR(VLOOKUP(E:E,Pemetaan!A:C,2,FALSE),"?"))))',
  '=ARRAYFORMULA(IF(ROW(E:E)=1,"Kategori",IF(E:E="","",IFERROR(VLOOKUP(E:E,Pemetaan!A:C,3,FALSE),"?"))))'])
arows = sorted(AM.items())
for (m, area, lok, name, case), n in arows:
    ws.append([m, area or None, lok or None, name, case, n])
for c in ws[1][:6]: c.font = HF; c.fill = HFILL
ws.freeze_panes = 'A2'; widths(ws, [7, 13, 22, 30, 22, 9, 12, 22, 19])
n_am = len(arows)

# ===== Sheet agregat =====
for name, f, w in [
    (AL, f'=QUERY({DE}!$A:$I,"select A, D, G, I, sum(F) where A is not null and I matches \'{KAT3}\' group by A, D, G, I label A \'Bulan\', D \'Lokasi\', G \'Regional\', I \'Kategori\', sum(F) \'Jumlah\'",1)', [7, 24, 13, 20, 10]),
    (AP, f'=QUERY({DE}!$A:$I,"select A, H, sum(F) where A is not null group by A, H label A \'Bulan\', H \'Parameter\', sum(F) \'Jumlah\'",1)', [7, 24, 10]),
    (AH, f'=QUERY({DE}!$A:$I,"select B, G, I, sum(F) where B is not null and I matches \'{KAT3}\' group by B, G, I label B \'Tanggal\', G \'Regional\', I \'Kategori\', sum(F) \'Jumlah\'",1)', [12, 13, 20, 10]),
    (AA, f'=QUERY({DA}!$A:$I,"select A, D, C, G, I, sum(F) where A is not null and I matches \'{KAT3}\' group by A, D, C, G, I label A \'Bulan\', D \'Nama AMT\', C \'Lokasi\', G \'Regional\', I \'Kategori\', sum(F) \'Jumlah\'",1)', [7, 30, 22, 13, 20, 10])]:
    ws = wb.create_sheet(name); ws['A1'] = f; widths(ws, w)

# ===== Parameter_Bulanan =====
ws = wb.create_sheet('Parameter_Bulanan', 1)
ws.append(['Parameter per Kategori per Bulan'])
ws.append(['Rumus: SUMIFS dari Data_Event (Bulan & Parameter). Bulan tanpa data dibiarkan kosong.'])
ws.append(['Baris data per bulan'] + [f'={cnt(m)}' for m in range(1, 13)])
ws.append(['Parameter'] + BLN + ['Total'])
ws['A1'].font = H1; ws['A2'].font = NOTE
for c in ws[3]: c.font = HELP
head(ws, 4)
def pcell(r, m):
    col = CL(m + 1); base = sumifs("A", m, "B", f"$A{r}", col='C', sh=AP)
    return f'=IF({col}$3=0,"",{base})'
r = 5; totrows = []
for k in KATS:
    ws.append([k.upper()]); ws.cell(r, 1).font = BOLD; r += 1
    start = r
    for l, kk, _ in PAR:
        if kk != k: continue
        ws.append([l] + [pcell(r, m) for m in range(1, 13)] + [f'=SUM(B{r}:M{r})']); r += 1
    ws.append([f'TOTAL {k}'] + [f'=IF(COUNT({CL(m+1)}{start}:{CL(m+1)}{r-1})=0,"",SUM({CL(m+1)}{start}:{CL(m+1)}{r-1}))' for m in range(1, 13)] + [f'=SUM(B{r}:M{r})'])
    for c in ws[r]: c.font = BOLD; c.fill = TOT
    totrows.append(r); r += 1
    ws.append([]); r += 1
ws.append(['TOTAL 3 KATEGORI'] + [f'=IF(COUNT({",".join(CL(m+1)+str(t) for t in totrows)})=0,"",{"+".join("N("+CL(m+1)+str(t)+")" for t in totrows)})' for m in range(1, 13)] + [f'=SUM(B{r}:M{r})'])
for c in ws[r]: c.font = BOLD; c.fill = TOT
TOT3ROW = r; r += 1
ws.append([]); ws.append([]); r += 2
ws.append(['PARAMETER LAIN (di luar 3 kategori, tidak masuk total)']); ws.cell(r, 1).font = BOLD; r += 1
ws.append(['Parameter'] + BLN + ['Total']); head(ws, r); r += 1
for l, kk, _ in PAR:
    if kk == 'Parameter lain':
        ws.append([l] + [pcell(r, m) for m in range(1, 13)] + [f'=SUM(B{r}:M{r})']); r += 1
widths(ws, [28] + [11] * 13); ws.freeze_panes = 'B5'

# ===== Rekap_Lokasi_YTD (dibuat dulu; dirujuk Rekap_Regional) =====
ALLLOK = sorted({(regrow(a, l), l) for (m, dt, a, l, c) in EV if l and CASE[c][1] in KATS and regrow(a, l) != '?'},
                key=lambda x: (REG.index(x[0].upper()), x[1]))
lok_tot = collections.Counter(); lok_m = collections.defaultdict(collections.Counter)
for (m, dt, a, l, c), n in EV.items():
    if CASE[c][1] in KATS and l: lok_tot[(regrow(a, l), l)] += n; lok_m[m][(regrow(a, l), l)] += n

def lokasi_sheet(name, m, pm, title):
    ws = wb.create_sheet(name)
    ws['A1'] = title; ws['A1'].font = H1
    if m:
        ws['I1'] = 'Bulan →'; ws['J1'] = m; ws['K1'] = pm if pm else None; ws['L1'] = '← bulan pembanding'
        for c in ('J1', 'K1'): ws[c].fill = INPUT; ws[c].font = BOLD
        ws['A2'] = 'Rumus: SUMIFS dari Data_Event (Bulan di J1, Lokasi, Regional, Kategori). Sel kuning J1/K1 bisa diganti untuk bulan lain. Lokasi baru: salin satu baris lalu ganti nama lokasi & regional.'
    else:
        ws['A2'] = 'Rumus: SUMIFS dari Data_Event untuk semua bulan yang ada di Data_Event. Rata-rata/bulan = total ÷ bulan terdata.'
    ws['A2'].font = NOTE
    hdr = ['Lokasi', 'Regional'] + KATS + ['Total']
    if m: hdr += (['=IF(K1="","Total bln lalu","Total "&CHOOSE(K1,' + ','.join(f'"{b}"' for b in B3) + '))', 'Kenaikan'] if pm else [])
    else: hdr += ['Rata-rata/bulan', 'Bulan Terdata']
    ws.append([]); ws.append(hdr); head(ws, 4)
    order = sorted(ALLLOK, key=lambda x: (-(lok_m[m][x] if m else lok_tot[x]), REG.index(x[0].upper()), x[1]))
    for i, (rg, l) in enumerate(order):
        r = 5 + i
        crit = ['D', f'$A{r}', 'G', f'$B{r}'] + (['A', '$J$1'] if m else [])
        row = [l, rg] + [f'={sl(*crit, "I", CL(3 + j) + "$4")}' for j in range(3)] + [f'=SUM(C{r}:E{r})']
        if m and pm:
            row += [f'=IF($K$1="","",IF({cntl("$K$1")}=0,"",{sl("D", f"$A{r}", "G", f"$B{r}", "A", "$K$1")}))', growth(f'F{r}', f'G{r}')]
        elif not m:
            row += [f'=IF(H{r}=0,"",F{r}/H{r})', f'=COUNTUNIQUEIFS({AL}!$A:$A,{AL}!$B:$B,$A{r},{AL}!$C:$C,$B{r},{AL}!$E:$E,">0")']
        ws.append(row)
        if m and pm: ws.cell(r, 8).number_format = '0.0%'
        if not m: ws.cell(r, 7).number_format = '#,##0.0'
    widths(ws, [24, 13, 17, 17, 19, 11, 14, 11]); ws.freeze_panes = 'A5'
    return ws

# ===== Rekap_Regional =====
ws = wb.create_sheet('Rekap_Regional', 2)
ws['A1'] = 'Rekap per Regional'; ws['A1'].font = H1
ws['A2'] = 'Rumus: SUMIFS dari Data_Event (Bulan, Regional, Kategori). Tiap blok punya angka bulan di sel kuning kolom J (bulan) dan K (pembanding).'; ws['A2'].font = NOTE
r = 4
for i, m in enumerate(MONTHS):
    pm = (m - 1) if m > 1 else None
    ws.cell(r, 1, f'Rekap per Regional – {BLN[m-1]} 2026' + (f' vs {BLN[pm-1]} 2026' if pm else '')).font = BOLD
    ws.cell(r, 9, 'Bulan →').font = HELP; ws.cell(r, 10, m); ws.cell(r, 11, pm)
    for c in (10, 11): ws.cell(r, c).fill = INPUT
    t = r; r += 1
    choose = lambda ref: 'CHOOSE(' + ref + ',' + ','.join(f'"{b}"' for b in BLN) + ')'
    hdr = ['Regional'] + KATS + [f'="Total "&{choose(f"$J${t}")}'] + ([f'="Total "&{choose(f"$K${t}")}', 'Kenaikan'] if pm else []) + ['Jumlah Lokasi']
    for j, v in enumerate(hdr, 1): ws.cell(r, j, v)
    head(ws, r); h = r; r += 1
    first = r
    for rg in REG:
        R = RN[rg]
        ws.cell(r, 1, R)
        for j in range(3): ws.cell(r, 2 + j, f'={sl("A", f"$J${t}", "G", f"$A{r}", "I", CL(2 + j) + f"${h}")}')
        ws.cell(r, 5, f'=SUM(B{r}:D{r})')
        c = 6
        if pm:
            ws.cell(r, 6, f'=IF({cntl(f"$K${t}")}=0,"",{sl("A", f"$K${t}", "G", f"$A{r}")})')
            ws.cell(r, 7, growth(f'E{r}', f'F{r}')); ws.cell(r, 7).number_format = '0.0%'; c = 8
        ws.cell(r, c, f'=COUNTIFS(Rekap_Lokasi_YTD!$B:$B,$A{r},Rekap_Lokasi_YTD!$F:$F,">0")')
        r += 1
    ws.cell(r, 1, 'TOTAL')
    for j in range(2, 6): ws.cell(r, j, f'=SUM({CL(j)}{first}:{CL(j)}{r-1})')
    if pm:
        ws.cell(r, 6, f'=SUM(F{first}:F{r-1})'); ws.cell(r, 7, growth(f'E{r}', f'F{r}')); ws.cell(r, 7).number_format = '0.0%'
        ws.cell(r, 8, f'=SUM(H{first}:H{r-1})')
    else: ws.cell(r, 6, f'=SUM(F{first}:F{r-1})')
    for cc in ws[r][:8]: cc.font = BOLD; cc.fill = TOT
    r += 2
ws.cell(r, 1, 'Total event 3 kategori per regional per bulan').font = BOLD; r += 1
ws.cell(r, 1, 'Regional')
for m in range(1, 13): ws.cell(r, 1 + m, BLN[m - 1])
head(ws, r); r += 1
for rg in REG:
    R = RN[rg]; ws.cell(r, 1, R)
    for m in range(1, 13):
        ws.cell(r, 1 + m, f'=IF({cntl(m)}=0,"",{sl("A", m, "G", f"$A{r}")})')
    r += 1
widths(ws, [22, 17, 17, 19, 15, 15, 11, 13, 8, 7, 7, 11, 11])

# ===== Rekap_Lokasi_* =====
for m in MONTHS:
    pm = (m - 1) if m > 1 else None
    lokasi_sheet(f'Rekap_Lokasi_{B3[m-1]}', m, pm, f'Rekap per Lokasi – {BLN[m-1]} 2026')
lokasi_sheet('Rekap_Lokasi_YTD', None, None, 'Rekap per Lokasi – seluruh bulan di Data_Event')

# ===== Top_AMT =====
ws = wb.create_sheet('Top_AMT')
ws['A1'] = 'TOP 10 Pelanggaran AMT per Regional (per bulan dan YTD)'; ws['A1'].font = H1
ws['A2'] = 'Rumus: QUERY dari Data_AMT (kolom M:O, tersembunyi warna abu) mengambil 10 AMT teratas per regional; kolom E–G = SUMIFS per kategori; I = parameter terbanyak. Kolom kuning K (bulan) & L (label) bisa diganti.'; ws['A2'].font = NOTE
hdr = ['Bulan', 'Nama AMT', 'Regional', 'Lokasi'] + KATS + ['Total', 'Pelanggaran Terbanyak', None, 'Bulan (angka)', 'Label', 'q: Nama', 'q: Lokasi', 'q: Total']
for j, v in enumerate(hdr, 1): ws.cell(4, j, v)
head(ws, 4)
r = 5
blocks = [(m, BLN[m-1]) for m in AMONTHS] + [('YTD', 'YTD')]
for m, label in blocks:
    for rg in REG:
        R = RN[rg]
        cond = (f'A="&$K{r}&" and ' if m != 'YTD' else '')
        ws.cell(r, 13, f'=IFERROR(QUERY({AA}!$A:$F,"select B, C, sum(F) where {cond}D=\'"&$C{r}&"\' group by B, C order by sum(F) desc limit 10 label sum(F) \'\'",0),"")')
        for i in range(10):
            rr = r + i
            ws.cell(rr, 11, m if m != 'YTD' else 'YTD'); ws.cell(rr, 12, label)
            for c in (11, 12): ws.cell(rr, c).fill = INPUT
            ws.cell(rr, 1, f'=IF($M{rr}="","",$L{rr})')
            ws.cell(rr, 2, f'=IF($M{rr}="","",$M{rr})')
            ws.cell(rr, 3, R)
            ws.cell(rr, 4, f'=IF($M{rr}="","",$N{rr})')
            mc = ([ 'A', f'$K{rr}'] if m != 'YTD' else [])
            for j, k in enumerate(KATS):
                ws.cell(rr, 5 + j, f'=IF($M{rr}="","",{sa("D", f"$M{rr}", "C", f"$N{rr}", "G", f"$C{rr}", *mc, "I", chr(34) + k + chr(34))})')
            ws.cell(rr, 8, f'=IF($M{rr}="","",$O{rr})')
            cond2 = (f' and A="&$K{rr}&"' if m != 'YTD' else '')
            ws.cell(rr, 9, f'=IF($M{rr}="","",IFERROR(INDEX(QUERY({DA}!$A:$I,"select H, sum(F) where D="""&$M{rr}&""" and C="""&$N{rr}&""" and G=\'"&$C{rr}&"\'{cond2} and I matches \'{KAT3}\' group by H order by sum(F) desc limit 1 label sum(F) \'\'",0),1,1),""))')
            for c in (13, 14, 15): ws.cell(rr, c).font = Font(color='999999')
        r += 10
widths(ws, [10, 30, 12, 22, 16, 16, 18, 9, 22, 2, 9, 10, 24, 20, 9]); ws.freeze_panes = 'A5'

# ===== Harian =====
ws = wb.create_sheet('Harian')
ws['A1'] = 'Event per Hari per Regional (3 kategori)'; ws['A1'].font = H1
ws['A2'] = 'Rumus: QUERY pivot dari Agg_Harian (ringkasan harian Data_Event) di sel A4; bertambah otomatis kalau Data_Event ditambah.'; ws['A2'].font = NOTE
ws['A4'] = f'=QUERY({AH}!$A:$D,"select A, B, sum(D) where A is not null group by A, B pivot C label A \'Tanggal\', B \'Regional\'",1)'
widths(ws, [12, 13, 18, 18, 20])

# ===== Cek =====
ws = wb.create_sheet('Cek', 1)
ws['A1'] = 'Pemeriksaan otomatis'; ws['A1'].font = H1
chk = [
 ('Baris Data_Event dengan Case belum terpetakan (Kategori "?")', f'=COUNTIF({DE}!$I:$I,"?")', 'harus 0 → tambahkan Case ke Pemetaan'),
 ('Baris Data_Event dengan Regional tidak dikenali ("?")', f'=COUNTIF({DE}!$G:$G,"?")', 'harus 0 → tambahkan Area/Lokasi ke Pemetaan'),
 ('Baris Data_AMT dengan Case belum terpetakan', f'=COUNTIF({DA}!$I:$I,"?")', 'harus 0'),
 ('Jumlah lokasi (3 kategori) di Data_Event', f'=COUNTUNIQUEIFS({DE}!$D:$D,{DE}!$I:$I,"<>Tidak dihitung",{DE}!$I:$I,"<>Parameter lain")', ''),
 ('Jumlah lokasi di Rekap_Lokasi_YTD', '=COUNTA(Rekap_Lokasi_YTD!$A:$A)-3', 'harus sama dengan baris di atas → kalau kurang, tambahkan baris lokasi baru'),
 ('Total event 3 kategori di Data_Event', f'=SUMIFS({DE}!$F:$F,{DE}!$I:$I,"Driving Behaviour")+SUMIFS({DE}!$F:$F,{DE}!$I:$I,"Driver Discipline")+SUMIFS({DE}!$F:$F,{DE}!$I:$I,"Fatigue Management")', ''),
 ('Total 3 kategori di Rekap_Lokasi_YTD', '=SUM(Rekap_Lokasi_YTD!$F:$F)', 'selisih = event tanpa lokasi/regional'),
 ('Total 3 kategori di Agg_Lokasi', f'=SUM({AL}!$E:$E)', 'harus sama dengan total di Data_Event'),
]
ws.append([]); ws.append(['Pemeriksaan', 'Nilai', 'Keterangan']); head(ws, 3)
for t, f, k in chk: ws.append([t, f, k])
widths(ws, [60, 14, 70])

# ===== Kelengkapan =====
import calendar
ws = wb.create_sheet('Kelengkapan')
ws['A1'] = 'Kelengkapan data per bulan & regional (dari file mentah)'; ws['A1'].font = H1
ws['A2'] = 'Hari kosong = tidak ada event sama sekali. Hari rendah = di bawah 35% median harian regional tsb (hari Minggu/libur bisa wajar rendah). "Batas Excel" = sheet terisi 1.048.575 baris, kemungkinan ada baris yang terpotong.'; ws['A2'].font = NOTE
ws.append([]); ws.append(['Bulan', 'Regional', 'File sumber', 'Baris dihitung', 'Tanggal pertama', 'Tanggal terakhir', 'Hari kosong', 'Hari rendah', 'Catatan']); head(ws, 4)
LASTDAY = max(datetime.date.fromisoformat(r[0]) for d in INFO for r in d['DL'] if r[0])
GAPS = []
for d in sorted(INFO, key=lambda d: (d['mon'], REG.index(d['reg']) if d['reg'] in REG else 9)):
    m = d['mon']; days = collections.Counter()
    for ds, a_, l_, c_, n in d['DL']:
        if ds and CASE.get(c_, ('', ''))[1] != 'Tidak dihitung': days[int(ds[8:10])] += n
    nd = calendar.monthrange(2026, m)[1]
    if m == LASTDAY.month: nd = LASTDAY.day
    vals = sorted(days.values()); med = vals[len(vals)//2] if vals else 0
    miss = [x for x in range(1, nd + 1) if days[x] == 0]; low = [x for x in range(1, nd + 1) if 0 < days[x] < med * 0.35]
    sheets = [(f['file'].split('/')[-1], sh, nn) for f in d['info']['files'] for sh, nn in f['sheets'] if isinstance(nn, int) and nn > 0]
    perfile = collections.Counter(fn for fn, sh, nn in sheets)
    cap = [f'{fn} [{sh}]' for fn, sh, nn in sheets if nn >= 1048575 - 1 and perfile[fn] == 1]
    notes = []
    if cap: notes.append('Batas Excel: ' + '; '.join(cap))
    if d['info']['dup_removed']: notes.append(f"{d['info']['dup_removed']:,} baris dobel antar file dihapus".replace(',', '.'))
    om = sum(d['info']['outmonth'].values())
    if om: notes.append(f'{om:,} baris bertanggal di luar bulan tidak dihitung'.replace(',', '.'))
    ds_ = sorted(days)
    ws.append([BLN[m-1], RN.get(d['reg'], d['reg']), '\n'.join(sorted({f[0] for f in sheets})), d['info']['rows'],
               f'2026-{m:02d}-{ds_[0]:02d}' if ds_ else '', f'2026-{m:02d}-{ds_[-1]:02d}' if ds_ else '',
               ', '.join(map(str, miss)), ', '.join(map(str, low)), '; '.join(notes)])
    ws.cell(ws.max_row, 3).alignment = Alignment(wrap_text=True, vertical='top')
    ws.cell(ws.max_row, 4).number_format = '#,##0'
    R_ = RN.get(d['reg'], d['reg'])
    if m == LASTDAY.month:
        continue
    if len(miss) >= 3: GAPS.append(f'{BLN[m-1]} {R_}: tidak ada data tanggal {", ".join(map(str, miss))}.')
    elif miss: GAPS.append(f'{BLN[m-1]} {R_}: tanggal {", ".join(map(str, miss))} kosong.')
    if cap: GAPS.append(f'{BLN[m-1]} {R_}: sheet di file sumber terisi penuh sampai batas Excel 1.048.575 baris, kemungkinan ada event yang terpotong.')
GAPS.append(f'{BLN[LASTDAY.month-1]} masih berjalan: data s/d {LASTDAY.day} {BLN[LASTDAY.month-1]} dan tidak semua regional sudah sampai tanggal itu (lihat sheet Kelengkapan).')
widths(ws, [11, 12, 48, 14, 14, 14, 30, 30, 60]); ws.freeze_panes = 'A5'

# ===== Catatan (dibaca dashboard) =====
ws = wb.create_sheet('Catatan')
cat = ['PEMETAAN',
 'Semua angka dihitung dengan rumus dari sheet Data_Event dan Data_AMT; pemetaan Case → parameter ada di sheet Pemetaan.',
 'Driving Behaviour: Over Speed=OVERSPEED, Harsh Turn/Cornering=HARSH TURN, Harsh Braking=HARSH BREAKING, Harsh Acceleration=HARSH ACCELERATION, Driving > 4 Hours=DRIVING > 4 HOURS.',
 'Driver Discipline: Black Zone=BLACKZONE, Idling=IDLE, Menggunakan Telepon=PHONE DETECTION, Merokok/Vape=SMOKING DETECTION.',
 'Fatigue Management: Microsleep=DRIVER FATIGUE (asumsi, mohon konfirmasi), Menguap=YAWNING DETECTION.',
 'Parameter lain: Distraction Pengemudi=DRIVER DISTRACTION, Camera Covering=CAMERA COVERING ALARM. Tidak dihitung: SEAT BELT DETECTION, ILLEGAL SHUTDOWN, REST AREA, VIDEO LOSS/HIGH SPEED/ABNORMAL STORAGE ALARM.',
 'TOP AMT: nama dari kolom AMT 1 (NIP dalam kurung dibuang agar satu orang tidak terhitung dua kali); event tanpa nama AMT tidak ikut peringkat.',
 '', 'KETERBATASAN DATA'] + [f'{i}. {t}' for i, t in enumerate(GAPS, 1)]
for c in cat: ws.append([c])
ws['A1'].font = BOLD; ws['A9'].font = BOLD; widths(ws, [150])

order = ['Panduan', 'Cek', 'Parameter_Bulanan', 'Rekap_Regional'] + [f'Rekap_Lokasi_{B3[m-1]}' for m in MONTHS] + ['Rekap_Lokasi_YTD', 'Top_AMT', 'Harian', 'Catatan', 'Kelengkapan', 'Pemetaan', AL, AP, AH, AA, DE, DA]
wb._sheets = [wb[n] for n in order]
for n in ('Kelengkapan', 'Pemetaan', AL, AP, AH, AA, DE, DA): wb[n].sheet_properties.tabColor = '999999'
for n in ('Panduan', 'Cek'): wb[n].sheet_properties.tabColor = 'E8833A'
wb.save(OUTX)
print('Data_Event', n_ev, 'Data_AMT', n_am, 'lokasi', len(ALLLOK), 'months', MONTHS, AMONTHS, 'drop', dict(drop))
