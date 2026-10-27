# Agregat file mentah satu bulan (keluaran agregasi_jam.py) -> "Data <nn> <Bulan> 2026.xlsx" untuk folder "3. Data Dashboard"
# (sheet Event: Tanggal | Area | Lokasi | Case | Jumlah, sheet AMT: Area | Lokasi | Nama AMT | Case | Jumlah; lihat dataUtama_ di Kode.gs).
# Pemakaian: python3 -I bangun_data_bulan.py <bulan 1-12> <folder agregat> <folder keluaran> "Pengaturan Dashboard.xlsx" [sebelum yyyy-mm-dd]
#   sebelum: hanya tanggal sebelum tanggal ini (tanggal itu dan sesudahnya diambil dari file tarikan, lihat CFG.MULAI_TARIKAN).
import sys, json, glob, collections, openpyxl
from openpyxl.styles import Font, PatternFill
M = int(sys.argv[1]); AGG = sys.argv[2]; OUT = sys.argv[3]; ATUR = sys.argv[4]; SEBELUM = sys.argv[5] if len(sys.argv) > 5 else '9999'
BLN = ['Januari', 'Februari', 'Maret', 'April', 'Mei', 'Juni', 'Juli', 'Agustus', 'September', 'Oktober', 'November', 'Desember']
N = lambda x: ' '.join(str(x).replace('\xa0', ' ').split()).upper()
kenal = set()
for r in openpyxl.load_workbook(ATUR, read_only=True)['Pemetaan'].iter_rows(min_row=2, values_only=True):
    if r[0]: kenal.add(N(r[0]))
EV = collections.Counter(); AM = collections.Counter(); asing = collections.Counter(); hari = collections.defaultdict(set)
for f in sorted(glob.glob(f'{AGG}/m{M:02d}_*.json')):
    d = json.load(open(f))
    for ds, h, a, l, c, n in d['EJ']:
        if not ds or ds >= SEBELUM: continue
        if N(c) not in kenal: asing[N(c)] += n
        EV[(ds, a, l, c)] += n; hari[d['reg']].add(ds)
    for ds, h, a, l, nm, c, n in d['AJ']:
        if not ds or ds >= SEBELUM: continue
        AM[(a, l, nm, c)] += n
HF = Font(bold=True, color='FFFFFF'); HFILL = PatternFill('solid', fgColor='4A7A53')
wb = openpyxl.Workbook(); se = wb.active; se.title = 'Event'; sa = wb.create_sheet('AMT')
for sh, hdr, rows in ((se, ['Tanggal', 'Area', 'Lokasi', 'Case', 'Jumlah'], sorted(EV.items())), (sa, ['Area', 'Lokasi', 'Nama AMT', 'Case', 'Jumlah'], sorted(AM.items()))):
    sh.append(hdr)
    for c in sh[1]: c.font = HF; c.fill = HFILL
    for k, n in rows: sh.append(list(k) + [n])
    sh.freeze_panes = 'A2'
se.column_dimensions['A'].number_format = '@'
p = f'{OUT}/Data {M:02d} {BLN[M-1]} 2026.xlsx'
wb.save(p)
print(p, 'Event', len(EV), 'baris', sum(EV.values()), 'event; AMT', len(AM), 'baris')
for r in sorted(hari): print(' ', r, min(hari[r]), '..', max(hari[r]), len(hari[r]), 'hari')
if asing: print('Case belum ada di Pemetaan:', asing.most_common())
