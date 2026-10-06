# Ekspor data dashboard ke file per bulan untuk folder Google Drive "3. Data Dashboard".
# Pemakaian: python3 -I ekspor_drive.py <summary_rumus.xlsx> <folder_keluaran>
# Masukan = xlsx hasil bangun_summary_rumus.py (sheet Data_Event, Data_AMT, Pemetaan, Catatan, Kelengkapan).
import sys, os, collections, datetime, openpyxl
from openpyxl.styles import Font, PatternFill
SRC, OUT = sys.argv[1], sys.argv[2]
BLN = ['Januari', 'Februari', 'Maret', 'April', 'Mei', 'Juni', 'Juli', 'Agustus', 'September', 'Oktober', 'November', 'Desember']
HF = Font(bold=True, color='FFFFFF'); HFILL = PatternFill('solid', fgColor='4A7A53')
wb = openpyxl.load_workbook(SRC, read_only=True)
ev = collections.defaultdict(list); am = collections.defaultdict(list)
for r in wb['Data_Event'].iter_rows(min_row=2, max_col=6, values_only=True):
    if not r[4]: continue
    d = r[1]; ds = d.strftime('%Y-%m-%d') if isinstance(d, (datetime.date, datetime.datetime)) else str(d)[:10]
    ev[int(r[0])].append([ds, r[2], r[3], r[4], r[5]])
for r in wb['Data_AMT'].iter_rows(min_row=2, max_col=6, values_only=True):
    if not r[4]: continue
    am[int(r[0])].append([r[1], r[2], r[3], r[4], r[5]])
def sheet(w, name, hdr, rows, widths):
    ws = w.create_sheet(name); ws.append(hdr)
    for c in ws[1]: c.font = HF; c.fill = HFILL
    for x in rows: ws.append(x)
    for i, wd in enumerate(widths): ws.column_dimensions[chr(65 + i)].width = wd
    ws.freeze_panes = 'A2'
for m in sorted(set(ev) | set(am)):
    w = openpyxl.Workbook(); w.remove(w.active)
    sheet(w, 'Event', ['Tanggal', 'Area', 'Lokasi', 'Case', 'Jumlah'], ev[m], [12, 14, 26, 26, 9])
    sheet(w, 'AMT', ['Area', 'Lokasi', 'Nama AMT', 'Case', 'Jumlah'], am[m], [14, 26, 32, 26, 9])
    p = os.path.join(OUT, f'Data {m:02d} {BLN[m-1]} 2026.xlsx'); w.save(p)
    print(p, len(ev[m]), len(am[m]), os.path.getsize(p))
# Pengaturan: Pemetaan + Catatan (+ Kelengkapan)
w = openpyxl.Workbook(); w.remove(w.active)
for name in ('Pemetaan', 'Catatan', 'Kelengkapan'):
    ws = w.create_sheet(name)
    rows = [list(r) for r in wb[name].iter_rows(values_only=True)]
    if name == 'Pemetaan':   # Case di data yang sengaja tidak dihitung -> dicatat supaya dashboard tidak memberi peringatan
        known = {str(r[0]).strip().upper() for r in rows[1:] if r[0]}
        allc = sorted({str(x[3]).strip().upper() for v in ev.values() for x in v} | {str(x[3]).strip().upper() for v in am.values() for x in v})
        extra = [c for c in allc if c and c not in known]
        i = next((j for j, r in enumerate(rows) if j and not r[0]), len(rows))
        for c in extra:
            if i >= len(rows): rows.append([None] * 9)
            rows[i][0:3] = [c, '(tidak dihitung)', 'Tidak dihitung']; i += 1
        print('case tidak dihitung:', extra)
    if name == 'Catatan':
        rows = [[('Semua angka dihitung oleh dashboard dari file Data per bulan di folder 3. Data Dashboard; pemetaan Case → parameter ada di sheet Pemetaan file Pengaturan Dashboard.' if 'dihitung dengan rumus' in str(r[0]) else r[0])] + list(r[1:]) for r in rows]
    for r in rows: ws.append(r)
    if name == 'Pemetaan':
        for c in ws[1]:
            if c.value: c.font = HF; c.fill = HFILL
        for col, wd in zip('ABCDEFGHI', [30, 24, 20, 2, 14, 13, 2, 26, 13]): ws.column_dimensions[col].width = wd
w['Catatan'].column_dimensions['A'].width = 140
p = os.path.join(OUT, 'Pengaturan Dashboard.xlsx'); w.save(p); print(p, os.path.getsize(p))
